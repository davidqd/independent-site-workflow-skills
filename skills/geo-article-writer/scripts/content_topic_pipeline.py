#!/usr/bin/env python3
"""Create, validate, claim, and update an evidence-backed article topic queue."""

from __future__ import annotations

import argparse
import json
import os
import random
import re
import unicodedata
import uuid
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = [
    "选题ID",
    "文章主题",
    "主关键词",
    "次关键词",
    "搜索意图",
    "AIDA阶段",
    "目标受众",
    "目标市场",
    "市场定位属性",
    "内容类型",
    "内容集群",
    "核心用户问题",
    "生成渠道",
    "生成依据",
    "证据URL或查询",
    "关键词数据来源",
    "搜索量",
    "Ads竞争度",
    "SEO难度",
    "趋势",
    "需求得分",
    "业务匹配得分",
    "意图匹配得分",
    "受众匹配得分",
    "内容缺口得分",
    "证据可信得分",
    "优先级总分",
    "建议内部链接",
    "建议CTA",
    "建议字数",
    "处理状态",
    "认领时间",
    "完成日期",
    "内容链接",
    "WordPress 内容ID",
    "最终标题",
    "错误说明",
    "产品ID",
    "产品名称",
    "产品关键词",
    "产品库文件",
    "产品库工作表",
    "产品库行号",
    "产品参数依据",
]

ALIASES = {
    "选题ID": ["选题ID", "Topic ID", "topic_id"],
    "文章主题": ["文章主题", "内容主题", "主题", "Article Topic", "Topic", "Title", "title"],
    "主关键词": ["主关键词", "Primary Keyword", "primary_keyword"],
    "次关键词": ["次关键词", "Secondary Keywords", "secondary_keywords"],
    "搜索意图": ["搜索意图", "Search Intent", "search_intent"],
    "AIDA阶段": ["AIDA阶段", "AIDA Stage", "aida_stage"],
    "目标受众": ["目标受众", "Target Audience", "target_audience"],
    "目标市场": ["目标市场", "Target Market", "target_market"],
    "市场定位属性": ["市场定位属性", "Market Positioning", "market_positioning"],
    "内容类型": ["内容类型", "Content Type", "content_type"],
    "内容集群": ["内容集群", "Content Cluster", "content_cluster"],
    "核心用户问题": ["核心用户问题", "Core User Question", "core_user_question"],
    "生成渠道": ["生成渠道", "Source Channels", "source_channels"],
    "生成依据": ["生成依据", "Topic Rationale", "topic_rationale"],
    "证据URL或查询": ["证据URL或查询", "Evidence URL or Query", "evidence"],
    "关键词数据来源": ["关键词数据来源", "Keyword Data Source", "keyword_data_source"],
    "搜索量": ["搜索量", "Search Volume", "search_volume"],
    "Ads竞争度": ["Ads竞争度", "Ads Competition", "ads_competition"],
    "SEO难度": ["SEO难度", "SEO Difficulty", "seo_difficulty"],
    "趋势": ["趋势", "Trend", "trend"],
    "需求得分": ["需求得分", "Demand Score", "demand_score"],
    "业务匹配得分": ["业务匹配得分", "Business Fit Score", "business_fit_score"],
    "意图匹配得分": ["意图匹配得分", "Intent Fit Score", "intent_fit_score"],
    "受众匹配得分": ["受众匹配得分", "Audience Fit Score", "audience_fit_score"],
    "内容缺口得分": ["内容缺口得分", "Content Gap Score", "content_gap_score"],
    "证据可信得分": ["证据可信得分", "Evidence Confidence Score", "evidence_confidence_score"],
    "建议内部链接": ["建议内部链接", "Suggested Internal Link", "suggested_internal_link"],
    "建议CTA": ["建议CTA", "Suggested CTA", "suggested_cta"],
    "建议字数": ["建议字数", "Suggested Word Count", "suggested_word_count"],
    "处理状态": ["处理状态", "发布状态", "Status", "status"],
    "产品ID": ["产品ID", "产品编号", "SKU", "Product ID", "product_id", "sku"],
    "产品名称": ["产品名称", "产品名", "Product Name", "product_name"],
    "产品关键词": ["产品关键词", "原始产品关键词", "Product Keyword", "product_keyword"],
    "产品库文件": ["产品库文件", "Product Library File", "product_library_file"],
    "产品库工作表": ["产品库工作表", "Product Library Sheet", "product_library_sheet"],
    "产品库行号": ["产品库行号", "Product Library Row", "product_library_row"],
    "产品参数依据": ["产品参数依据", "Product Parameter Evidence", "product_parameter_evidence"],
}

REQUIRED_FIELDS = [
    "文章主题",
    "主关键词",
    "搜索意图",
    "AIDA阶段",
    "目标受众",
    "市场定位属性",
    "核心用户问题",
    "生成渠道",
    "生成依据",
    "证据URL或查询",
    "关键词数据来源",
]
SCORE_FIELDS = [
    "需求得分",
    "业务匹配得分",
    "意图匹配得分",
    "受众匹配得分",
    "内容缺口得分",
    "证据可信得分",
]
WEIGHTS = [0.20, 0.25, 0.20, 0.15, 0.10, 0.10]
AIDA_MAP = {
    "attention": "Attention",
    "注意": "Attention",
    "认知": "Attention",
    "interest": "Interest",
    "兴趣": "Interest",
    "desire": "Desire",
    "欲望": "Desire",
    "考虑": "Desire",
    "action": "Action",
    "行动": "Action",
    "转化": "Action",
}
INTENT_MAP = {
    "informational": "信息型",
    "information": "信息型",
    "信息型": "信息型",
    "commercial": "商业调研型",
    "commercial investigation": "商业调研型",
    "商业调研型": "商业调研型",
    "transactional": "交易型",
    "交易型": "交易型",
    "navigational": "导航型",
    "导航型": "导航型",
}
DONE_VALUES = {"草稿已创建", "已完成", "已发布", "draft created", "finished", "published", "done", "complete", "completed"}
CLAIMED_VALUES = {"写作中", "claimed", "in progress", "in-progress"}
PRIORITY_VALUES = {"优先", "priority", "high-priority"}
RETRY_VALUES = {"失败待重试", "retry", "failed", "failed-retry"}


def text(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return "; ".join(str(item).strip() for item in value if str(item).strip())
    return str(value).strip()


def normalized_key(value):
    value = unicodedata.normalize("NFKC", text(value)).lower()
    return re.sub(r"[^\w\u4e00-\u9fff]+", "", value)


def lookup(item, field):
    for key in ALIASES.get(field, [field]):
        if key in item and item[key] is not None:
            return item[key]
    return ""


def normalize_aida(value):
    raw = text(value)
    return AIDA_MAP.get(raw.lower(), AIDA_MAP.get(raw, raw))


def normalize_intent(value):
    raw = text(value)
    return INTENT_MAP.get(raw.lower(), INTENT_MAP.get(raw, raw))


def parse_score(value):
    try:
        score = float(value)
    except (TypeError, ValueError):
        return None
    return score if 0 <= score <= 100 else None


def priority_score(record):
    values = [parse_score(record.get(field)) for field in SCORE_FIELDS]
    if any(value is None for value in values):
        return None
    return round(sum(value * weight for value, weight in zip(values, WEIGHTS)), 1)


def normalize_record(item, index):
    record = {field: lookup(item, field) for field in ALIASES}
    record["选题ID"] = text(record["选题ID"]) or f"T{index:04d}"
    record["文章主题"] = text(record["文章主题"])
    record["主关键词"] = text(record["主关键词"])
    record["次关键词"] = text(record["次关键词"])
    record["搜索意图"] = normalize_intent(record["搜索意图"])
    record["AIDA阶段"] = normalize_aida(record["AIDA阶段"])
    for field in [
        "目标受众", "目标市场", "市场定位属性", "内容类型", "内容集群",
        "核心用户问题", "生成渠道", "生成依据", "证据URL或查询",
        "关键词数据来源", "趋势", "建议内部链接", "建议CTA", "产品ID",
        "产品名称", "产品关键词", "产品库文件", "产品库工作表", "产品参数依据",
    ]:
        record[field] = text(record[field])
    for field in ["搜索量", "Ads竞争度", "SEO难度", "建议字数", "产品库行号", *SCORE_FIELDS]:
        value = record.get(field, "")
        record[field] = value if isinstance(value, (int, float)) else text(value)
    record["内容类型"] = record["内容类型"] or "博客文章"
    record["处理状态"] = text(record["处理状态"]) or "待处理"
    record["优先级总分"] = priority_score(record)
    return record


def split_channels(value):
    return [part.strip() for part in re.split(r"[;,，；|/]", text(value)) if part.strip()]


def validate_records(records, min_count, require_product_aida=False):
    errors = []
    warnings = []
    titles = {}
    keyword_intents = {}
    aida_counts = Counter()
    intent_counts = Counter()
    channels = set()
    product_pair_stages = {}

    if len(records) < min_count:
        errors.append(f"去重后只有 {len(records)} 条话题，要求至少 {min_count} 条")

    for number, record in enumerate(records, start=2):
        for field in REQUIRED_FIELDS:
            if not text(record.get(field)):
                errors.append(f"第 {number} 行缺少必填字段：{field}")
        aida = normalize_aida(record.get("AIDA阶段"))
        intent = normalize_intent(record.get("搜索意图"))
        if aida not in {"Attention", "Interest", "Desire", "Action"}:
            errors.append(f"第 {number} 行 AIDA 阶段无效：{text(record.get('AIDA阶段'))}")
        else:
            aida_counts[aida] += 1
        if intent not in {"信息型", "商业调研型", "交易型", "导航型"}:
            errors.append(f"第 {number} 行搜索意图无效：{text(record.get('搜索意图'))}")
        else:
            intent_counts[intent] += 1

        for field in SCORE_FIELDS:
            if parse_score(record.get(field)) is None:
                errors.append(f"第 {number} 行 {field} 必须是 0-100 的数字")

        title_key = normalized_key(record.get("文章主题"))
        if title_key:
            if title_key in titles:
                errors.append(f"第 {number} 行与第 {titles[title_key]} 行文章主题重复")
            else:
                titles[title_key] = number

        product_identity = normalized_key(record.get("产品ID") or record.get("产品名称"))
        product_keyword = normalized_key(record.get("产品关键词") or record.get("主关键词"))
        keyword_key = (
            (product_identity, product_keyword, intent, aida)
            if require_product_aida
            else (normalized_key(record.get("主关键词")), intent)
        )
        if keyword_key[0]:
            if keyword_key in keyword_intents:
                errors.append(
                    f"第 {number} 行与第 {keyword_intents[keyword_key]} 行主关键词和主要搜索意图重复，存在关键词蚕食风险"
                )
            else:
                keyword_intents[keyword_key] = number

        if require_product_aida:
            for field in ["产品名称", "产品关键词", "产品库文件", "产品库工作表", "产品库行号", "产品参数依据"]:
                if not text(record.get(field)):
                    errors.append(f"第 {number} 行产品库模式缺少必填字段：{field}")
            if product_identity and product_keyword and aida in {"Attention", "Interest", "Desire", "Action"}:
                pair = (product_identity, product_keyword)
                stages = product_pair_stages.setdefault(pair, {})
                if aida in stages:
                    errors.append(
                        f"第 {number} 行与第 {stages[aida]} 行属于同一产品关键词组合的重复 {aida} 主题"
                    )
                else:
                    stages[aida] = number

        channels.update(split_channels(record.get("生成渠道")))
        if not text(record.get("搜索量")):
            warnings.append(f"第 {number} 行搜索量为空，应填写真实数值或待查询")

    missing_aida = sorted({"Attention", "Interest", "Desire", "Action"} - set(aida_counts))
    if missing_aida:
        errors.append(f"AIDA 阶段覆盖不完整：缺少 {', '.join(missing_aida)}")
    if len(channels) < 4:
        errors.append(f"生成渠道只有 {len(channels)} 类，要求至少 4 类")
    channel_text = " ".join(channels).lower()
    if not any(token in channel_text for token in ["企业", "网站", "官网", "品牌", "site", "brand", "product", "service"]):
        errors.append("生成渠道缺少企业自身定位证据")
    if not any(token in channel_text for token in ["关键词", "keyword"]):
        errors.append("生成渠道缺少关键词数据")

    incomplete_product_pairs = []
    if require_product_aida:
        required_stages = {"Attention", "Interest", "Desire", "Action"}
        for pair, stages in sorted(product_pair_stages.items()):
            missing = sorted(required_stages - set(stages))
            if missing:
                incomplete_product_pairs.append({
                    "product": pair[0],
                    "product_keyword": pair[1],
                    "missing_stages": missing,
                })
        if incomplete_product_pairs:
            errors.append(f"有 {len(incomplete_product_pairs)} 个产品关键词组合未完整覆盖 A/I/D/A 四阶段")

    return {
        "valid": not errors,
        "topic_count": len(records),
        "min_count": min_count,
        "aida_distribution": dict(aida_counts),
        "intent_distribution": dict(intent_counts),
        "source_channels": sorted(channels),
        "product_aida_required": require_product_aida,
        "product_keyword_pair_count": len(product_pair_stages),
        "incomplete_product_pairs": incomplete_product_pairs,
        "errors": errors,
        "warnings": warnings,
    }


def load_json_records(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    items = data.get("topics") if isinstance(data, dict) else data
    if not isinstance(items, list):
        raise SystemExit("JSON 必须是数组，或包含 topics 数组的对象")
    return [normalize_record(item, index) for index, item in enumerate(items, start=1) if isinstance(item, dict)]


def header_map(ws):
    return {text(cell.value): cell.column for cell in ws[1] if cell.value is not None}


def ensure_column(ws, headers, name):
    if name in headers:
        return headers[name]
    column = ws.max_column + 1
    ws.cell(1, column).value = name
    headers[name] = column
    return column


def worksheet_records(ws):
    headers = header_map(ws)
    records = []
    for row in range(2, ws.max_row + 1):
        item = {name: ws.cell(row, column).value for name, column in headers.items()}
        if text(item.get("文章主题") or item.get("主题") or item.get("内容主题")):
            records.append(normalize_record(item, row - 1))
    return records


def atomic_save(wb, path):
    path = Path(path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        with temporary.open("wb") as handle:
            wb.save(handle)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def style_queue_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    widths = {
        "A": 12, "B": 42, "C": 24, "D": 32, "E": 15, "F": 14, "G": 24,
        "H": 18, "I": 34, "J": 16, "K": 22, "L": 36, "M": 28, "N": 44,
        "O": 42, "P": 28, "Q": 12, "R": 12, "S": 12, "T": 14,
    }
    for column in range(21, 28):
        widths[ws.cell(1, column).column_letter] = 14
    widths.update({"AB": 34, "AC": 24, "AD": 12, "AE": 16, "AF": 20, "AG": 20, "AH": 34, "AI": 20, "AJ": 34, "AK": 40})
    widths.update({"AL": 16, "AM": 28, "AN": 28, "AO": 42, "AP": 22, "AQ": 14, "AR": 52})
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 17).number_format = "#,##0"
        ws.cell(row, 18).number_format = "0.00"
        for column in range(21, 27):
            ws.cell(row, column).number_format = "0"
        ws.cell(row, 27).number_format = "0.0"
        ws.cell(row, 30).number_format = "#,##0"
        ws.cell(row, 32).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row, 33).number_format = "yyyy-mm-dd hh:mm"
    if ws.max_row >= 2:
        last_column = ws.cell(1, ws.max_column).column_letter
        table = Table(displayName="TopicQueueTable", ref=f"A1:{last_column}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)


def add_summary_sheet(wb, queue_title):
    ws = wb.create_sheet("选题统计")
    ws.append(["指标", "数量"])
    ws.append(["话题总数", f"=COUNTA('{queue_title}'!B:B)-1"])
    for stage in ["Attention", "Interest", "Desire", "Action"]:
        ws.append([stage, f'=COUNTIF(\'{queue_title}\'!F:F,"{stage}")'])
    for status in ["待处理", "优先", "写作中", "失败待重试", "草稿已创建", "已完成", "已发布"]:
        ws.append([status, f'=COUNTIF(\'{queue_title}\'!AE:AE,"{status}")'])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 16
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")


def create_queue(args):
    output = Path(args.file).expanduser()
    if output.exists() and not args.force:
        raise SystemExit(f"输出文件已存在：{output}；如确认覆盖，请添加 --force")
    records = load_json_records(args.input)
    report = validate_records(records, args.min_count, args.require_product_aida)
    if not report["valid"]:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        raise SystemExit(1)

    wb = Workbook()
    ws = wb.active
    ws.title = args.sheet
    ws.append(HEADERS)
    for row_number, record in enumerate(records, start=2):
        row = []
        for header in HEADERS:
            if header == "优先级总分":
                row.append(
                    f"=ROUND(U{row_number}*20%+V{row_number}*25%+W{row_number}*20%+X{row_number}*15%+Y{row_number}*10%+Z{row_number}*10%,1)"
                )
            elif header in {"认领时间", "完成日期", "内容链接", "WordPress 内容ID", "最终标题", "错误说明"}:
                row.append("")
            else:
                row.append(record.get(header, ""))
        ws.append(row)
    style_queue_sheet(ws)
    add_summary_sheet(wb, ws.title)
    atomic_save(wb, output)
    report["file"] = str(output.resolve())
    report["sheet"] = ws.title
    print(json.dumps(report, ensure_ascii=False, indent=2))


def open_queue(path, sheet=None):
    path = Path(path).expanduser().resolve()
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return path, wb, ws


def validate_queue(args):
    _, _, ws = open_queue(args.file, args.sheet)
    report = validate_records(worksheet_records(ws), args.min_count, args.require_product_aida)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if not report["valid"]:
        raise SystemExit(1)


def priority_from_row(ws, row, headers):
    values = []
    for field in SCORE_FIELDS:
        column = headers.get(field)
        values.append(parse_score(ws.cell(row, column).value) if column else None)
    if any(value is None for value in values):
        return 0.0
    return round(sum(value * weight for value, weight in zip(values, WEIGHTS)), 1)


def json_value(value):
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    return value


def select_topics(args):
    path, wb, ws = open_queue(args.file, args.sheet)
    headers = header_map(ws)
    topic_col = headers.get("文章主题") or headers.get("主题") or headers.get("内容主题")
    if not topic_col:
        raise SystemExit("未找到文章主题列")
    status_col = ensure_column(ws, headers, "处理状态")
    claim_col = ensure_column(ws, headers, "认领时间")
    candidates = []
    for row in range(2, ws.max_row + 1):
        title = text(ws.cell(row, topic_col).value)
        if not title:
            continue
        status = text(ws.cell(row, status_col).value).lower()
        if status in DONE_VALUES or status in CLAIMED_VALUES:
            continue
        priority = priority_from_row(ws, row, headers)
        status_rank = 0 if status in PRIORITY_VALUES else 1 if status in RETRY_VALUES else 2
        candidates.append((row, status_rank, priority))

    if args.strategy == "random":
        random.SystemRandom().shuffle(candidates)
    elif args.strategy == "order":
        candidates.sort(key=lambda item: item[0])
    else:
        candidates.sort(key=lambda item: (item[1], -item[2], item[0]))
    selected = candidates[: min(args.count, len(candidates))]

    if args.claim and selected:
        claimed_at = datetime.now()
        for row, _, _ in selected:
            ws.cell(row, status_col).value = "写作中"
            ws.cell(row, claim_col).value = claimed_at
            ws.cell(row, claim_col).number_format = "yyyy-mm-dd hh:mm"
        atomic_save(wb, path)

    output = []
    for row, _, priority in selected:
        item = {name: json_value(ws.cell(row, column).value) for name, column in headers.items()}
        item.update({"row": row, "topic": text(ws.cell(row, topic_col).value), "priority_score": priority})
        if args.claim:
            item["处理状态"] = "写作中"
        output.append(item)
    print(json.dumps(output, ensure_ascii=False, indent=2))


def mark_topic(args):
    path, wb, ws = open_queue(args.file, args.sheet)
    if args.row < 2 or args.row > ws.max_row:
        raise SystemExit(f"第 {args.row} 行超出工作表范围")
    headers = header_map(ws)
    columns = {
        name: ensure_column(ws, headers, name)
        for name in ["处理状态", "完成日期", "内容链接", "WordPress 内容ID", "最终标题", "错误说明"]
    }
    ws.cell(args.row, columns["处理状态"]).value = args.status
    if args.status in DONE_VALUES:
        ws.cell(args.row, columns["完成日期"]).value = datetime.now()
        ws.cell(args.row, columns["完成日期"]).number_format = "yyyy-mm-dd hh:mm"
    else:
        ws.cell(args.row, columns["完成日期"]).value = None
    if args.url:
        ws.cell(args.row, columns["内容链接"]).value = args.url
    if args.content_id:
        ws.cell(args.row, columns["WordPress 内容ID"]).value = str(args.content_id)
    if args.title:
        ws.cell(args.row, columns["最终标题"]).value = args.title
    if args.error:
        ws.cell(args.row, columns["错误说明"]).value = args.error
    elif args.status in DONE_VALUES:
        ws.cell(args.row, columns["错误说明"]).value = ""
    atomic_save(wb, path)

    _, _, verify_ws = open_queue(path, ws.title)
    verify_headers = header_map(verify_ws)
    actual_status = text(verify_ws.cell(args.row, verify_headers["处理状态"]).value)
    if actual_status != args.status:
        raise SystemExit(f"回写验证失败：期望状态 {args.status}，实际状态 {actual_status}")
    print(json.dumps({"updated": True, "row": args.row, "status": actual_status}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    create = sub.add_parser("create", help="Create a validated Excel queue from topic JSON")
    create.add_argument("--input", required=True, help="JSON array or object containing a topics array")
    create.add_argument("--file", required=True, help="Output .xlsx path")
    create.add_argument("--sheet", default="文章选题库")
    create.add_argument("--min-count", type=int, default=200)
    create.add_argument("--require-product-aida", action="store_true")
    create.add_argument("--force", action="store_true", help="Overwrite an existing output file")
    create.set_defaults(func=create_queue)

    validate = sub.add_parser("validate", help="Validate an existing topic queue")
    validate.add_argument("--file", required=True)
    validate.add_argument("--sheet")
    validate.add_argument("--min-count", type=int, default=200)
    validate.add_argument("--require-product-aida", action="store_true")
    validate.set_defaults(func=validate_queue)

    select = sub.add_parser("select", help="Select or claim pending topics")
    select.add_argument("--file", required=True)
    select.add_argument("--sheet")
    select.add_argument("--count", type=int, default=3)
    select.add_argument("--strategy", choices=["priority", "order", "random"], default="priority")
    select.add_argument("--claim", action="store_true")
    select.set_defaults(func=select_topics)

    mark = sub.add_parser("mark", help="Update one topic after writing or WordPress draft creation")
    mark.add_argument("--file", required=True)
    mark.add_argument("--sheet")
    mark.add_argument("--row", type=int, required=True)
    mark.add_argument("--status", required=True)
    mark.add_argument("--url")
    mark.add_argument("--content-id", "--post-id", dest="content_id")
    mark.add_argument("--title")
    mark.add_argument("--error")
    mark.set_defaults(func=mark_topic)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
