#!/usr/bin/env python3
"""Select and update article topics from an Excel queue."""

from __future__ import annotations

import argparse
import json
import random
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


TOPIC_HEADERS = [
    "文章主题",
    "内容主题",
    "主题",
    "英文文章主题",
    "English Article Topic",
    "Content Topic",
    "Article Topic",
    "Topic",
    "Title",
]
STATUS_HEADERS = ["处理状态", "发布状态", "Status", "status"]
DONE_VALUES = {"已完成", "已发布", "finished", "published", "done", "complete", "completed"}
PRIORITY_STATUS_MARKERS = ("优先", "priority", "high-priority")


def header_map(ws):
    return {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value is not None}


def find_col(headers, candidates):
    for name in candidates:
        if name in headers:
            return headers[name]
    return None


def ensure_col(ws, headers, name):
    if name in headers:
        return headers[name]
    col = ws.max_column + 1
    ws.cell(1, col).value = name
    headers[name] = col
    return col


def row_value(ws, row, col):
    value = ws.cell(row, col).value
    return "" if value is None else str(value).strip()


def open_sheet(path, sheet=None):
    wb = load_workbook(path)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return wb, ws


def select_topics(args):
    path = Path(args.file).expanduser()
    wb, ws = open_sheet(path, args.sheet)
    headers = header_map(ws)
    topic_col = find_col(headers, TOPIC_HEADERS)
    if topic_col is None:
        raise SystemExit(f"No topic column found. Expected one of: {', '.join(TOPIC_HEADERS)}")
    status_col = find_col(headers, STATUS_HEADERS)
    if status_col is None:
        status_col = ensure_col(ws, headers, "处理状态")

    priority_candidates = []
    candidates = []
    for row in range(2, ws.max_row + 1):
        topic = row_value(ws, row, topic_col)
        if not topic:
            continue
        status = row_value(ws, row, status_col).lower()
        if status in DONE_VALUES:
            continue
        item = {"row": row, "topic": topic, "status": row_value(ws, row, status_col)}
        for key, col in headers.items():
            if key not in item:
                item[key] = row_value(ws, row, col)
        candidates.append(item)
        if any(marker in status for marker in PRIORITY_STATUS_MARKERS):
            priority_candidates.append(item)

    count = min(args.count, len(candidates))
    if count and priority_candidates:
        priority_count = min(count, len(priority_candidates))
        selected = random.SystemRandom().sample(priority_candidates, priority_count)
        if priority_count < count:
            selected_rows = {item["row"] for item in selected}
            remaining = [item for item in candidates if item["row"] not in selected_rows]
            selected.extend(random.SystemRandom().sample(remaining, count - priority_count))
    else:
        selected = random.SystemRandom().sample(candidates, count) if count else []
    wb.save(path)
    print(json.dumps(selected, ensure_ascii=False, indent=2))


def mark_topic(args):
    path = Path(args.file).expanduser()
    wb, ws = open_sheet(path, args.sheet)
    headers = header_map(ws)
    status_col = find_col(headers, STATUS_HEADERS)
    if status_col is None:
        status_col = ensure_col(ws, headers, "处理状态")
    date_col = find_col(headers, ["完成日期", "发布日期", "Completion Date", "Published At"])
    if date_col is None:
        date_col = ensure_col(ws, headers, "完成日期")
    url_col = find_col(headers, ["内容链接", "文章链接", "Content URL", "URL"])
    if url_col is None:
        url_col = ensure_col(ws, headers, "内容链接")
    content_id_col = find_col(headers, ["内容ID", "Content ID", "WP Post ID"])
    if content_id_col is None:
        content_id_col = ensure_col(ws, headers, "内容ID")
    title_col = find_col(headers, ["最终标题", "已发布标题", "Final Title", "Published Title"])
    if title_col is None:
        title_col = ensure_col(ws, headers, "最终标题")

    row = args.row
    if row < 2 or row > ws.max_row:
        raise SystemExit(f"Row {row} is outside worksheet range")

    ws.cell(row, status_col).value = args.status
    ws.cell(row, date_col).value = args.date or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    if args.url:
        ws.cell(row, url_col).value = args.url
    if args.content_id:
        ws.cell(row, content_id_col).value = str(args.content_id)
    if args.title:
        ws.cell(row, title_col).value = args.title
    wb.save(path)
    print(json.dumps({"updated": True, "row": row, "status": args.status}, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="cmd", required=True)

    p_select = sub.add_parser("select", help="Randomly select unpublished topics")
    p_select.add_argument("--file", required=True)
    p_select.add_argument("--sheet")
    p_select.add_argument("--count", type=int, default=2)
    p_select.set_defaults(func=select_topics)

    p_mark = sub.add_parser("mark", help="Mark a topic row after writing or publishing")
    p_mark.add_argument("--file", required=True)
    p_mark.add_argument("--sheet")
    p_mark.add_argument("--row", type=int, required=True)
    p_mark.add_argument("--status", default="已发布")
    p_mark.add_argument("--url")
    p_mark.add_argument(
        "--content-id",
        "--post-id",
        dest="content_id",
        help="Published content ID; --post-id is retained as a compatibility alias",
    )
    p_mark.add_argument("--title")
    p_mark.add_argument("--date")
    p_mark.set_defaults(func=mark_topic)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
